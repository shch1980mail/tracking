"""Импорт Excel. Один ImportBatch на каждый лист. Основной лист читается по позиции колонок
(фиксированный маппинг), остальные — по заголовкам и настраиваемому маппингу (settings)."""
import io
from dataclasses import dataclass
from datetime import date, datetime, timezone
from typing import Callable

from openpyxl import load_workbook
from sqlalchemy import delete
from sqlalchemy.ext.asyncio import AsyncSession

from models.imports import ImportBatch, Setting
from models.sheets import InboundShipment, LocalShipment, MoscowShipment, TransitToTencard
from models.shipment import GpsLink, Position as PositionRow, Shipment
from services.gps_service import sync_link
from services.normalize import clean, is_url, norm_status, parse_driver, parse_vehicle, to_date
import re

MAIN_SHEET = "Исх.логистика"
COLS = ["excel_row_num", "week_number", "order_id", "direction", "address", "weight", "comment",
        "ship_date", "plan_arrival", "carrier", "vehicle", "driver", "rate", "rate_plan", "_f",
        "vehicle_type", "termologer", "notes", "gps", "status_raw", "location_raw"]

def _int(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None

def _float(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None

def parse_row(row: tuple) -> dict | None:  # основной лист — без изменений относительно ядра
    r = dict(zip(COLS, list(row) + [None] * len(COLS)))
    oid = clean(r["order_id"])
    if not oid:
        return None
    d = {"order_id": oid, "excel_row_num": _int(r["excel_row_num"]), "week_number": _int(r["week_number"]),
         "direction": (clean(r["direction"]) or "").upper() or None, "address": clean(r["address"]),
         "weight": _float(r["weight"]), "comment": clean(r["comment"]),
         "ship_date": to_date(r["ship_date"]), "plan_arrival": to_date(r["plan_arrival"]),
         "carrier": clean(r["carrier"]), "rate": _int(r["rate"]), "rate_plan": _int(r["rate_plan"]),
         "termologer": clean(r["termologer"]), "notes": clean(r["notes"]),
         "status_raw": clean(r["status_raw"]), "location_raw": clean(r["location_raw"])}
    d["status"] = norm_status(d["status_raw"], d["notes"])
    if d["ship_date"] and d["plan_arrival"]:
        d["transit_days"] = (d["plan_arrival"] - d["ship_date"]).days
    g = r["gps"]
    if is_url(g):
        d["gps_link"] = g.strip()
    elif clean(g):
        d["notes"] = " | ".join(x for x in (d["notes"], clean(g)) if x)
    d.update(parse_vehicle(r["vehicle"]))
    d.update(parse_driver(r["driver"]))
    vt = clean(r["vehicle_type"])
    if vt and not d.get("vehicle_type"):
        d["vehicle_type"] = vt
    return d

# ---------- сериализация snapshot ----------
def _ser(v):
    return v.isoformat() if isinstance(v, (date, datetime)) else v

def _deser(model, k, v):
    if v is None:
        return None
    t = type(model.__table__.columns[k].type).__name__
    return date.fromisoformat(v) if t == "Date" else datetime.fromisoformat(v) if t == "DateTime" else v

def _aware(dt):
    return dt if dt is None or dt.tzinfo else dt.replace(tzinfo=timezone.utc)

def _coerce(model, field, v):
    col = model.__table__.columns.get(field)
    if col is None:
        raise ValueError(f"неизвестное поле БД «{field}»")
    t = type(col.type).__name__
    if t == "Integer":
        return _int(v)
    if t == "Float":
        return _float(v)
    if t == "Date":
        d = to_date(v)
        if d is None and v not in (None, ""):
            raise ValueError(f"невалидная дата «{v}»")
        return d
    return clean(v)

async def _upsert(session, model, d, snap, errors, seen, rown):
    oid = d["order_id"]
    if oid in seen:
        errors.append(f"строка {rown}: дубль {oid} в файле (применена последняя запись)")
    seen.add(oid)
    obj = await session.get(model, oid)
    if obj is None:
        session.add(model(**d)); await session.flush()
        snap["created"].append(oid)
        return 1, 0
    prev, new = {}, {}
    for k, v in d.items():
        if k == "order_id" or (v is None and k != "gps_link"):
            continue
        cur = getattr(obj, k)
        if cur != v:
            prev[k], new[k] = _ser(cur), _ser(v)
            setattr(obj, k, v)
    if getattr(obj, "deleted_at", None):
        prev["deleted_at"], new["deleted_at"] = _ser(obj.deleted_at), None
        obj.deleted_at = None
    await session.flush()
    if prev:
        snap["updated"].append({"order_id": oid, "prev": prev, "new": new})
        return 0, 1
    return 0, 0

def _read(ws):
    it = ws.iter_rows(values_only=True)
    hdr = [str(h).strip().lower() if h is not None else "" for h in next(it, [])]
    return hdr, [(i, r) for i, r in enumerate(it, start=2) if any(c not in (None, "") for c in r)]

def _mapped(hdr, mapping):
    return {f: hdr.index(k.strip().lower()) for k, f in mapping.items() if k.strip().lower() in hdr}

# ---------- обработчики листов ----------
async def _handle_main(session, ws, mapping, snap, errors):
    ins = upd = 0; seen: set[str] = set()
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        try:
            d = parse_row(row)
            if d is None:
                continue
            a, b = await _upsert(session, Shipment, d, snap, errors, seen, i)
            ins, upd = ins + a, upd + b
            if d.get("gps_link"):
                await sync_link(session, d["order_id"], d["gps_link"])
        except Exception as e:
            errors.append(f"строка {i}: {e}")
    return ins, upd

FORECAST = re.compile(r"(\d{2})\.(\d{2})(?:\.(\d{4}))?")

def forecast_date(text, ship: date | None) -> date | None:
    m = FORECAST.search(text or "")
    if not m:
        return None
    day, mon = int(m.group(1)), int(m.group(2))
    year = int(m.group(3)) if m.group(3) else (ship.year + (1 if mon < ship.month else 0) if ship else date.today().year)
    try:
        return date(year, mon, day)
    except ValueError:
        return None

async def _handle_plan(session, ws, mapping, snap, errors):
    ins = upd = 0; seen: set[str] = set()
    hdr, rows = _read(ws); idx = _mapped(hdr, mapping)
    if "order_id" not in idx:
        raise ValueError("не найдена колонка «Заказ»")
    for i, row in rows:
        try:
            d = {f: _coerce(Shipment, f, row[j] if j < len(row) else None) for f, j in idx.items()}
            if not d.get("order_id"):
                errors.append(f"строка {i}: пустой order_id"); continue
            vt = d.get("vehicle_type")
            d.update(parse_vehicle(d.get("vehicle_raw"))); d.update(parse_driver(d.get("driver_raw")))
            if vt:
                d["vehicle_type"] = vt
            if (pa := forecast_date(d.get("comment"), d.get("ship_date"))):
                d["plan_arrival"] = pa
            d["status"] = "planned"
            ex = await session.get(Shipment, d["order_id"])
            if ex and ex.status == "delivered":
                d.pop("status")
            a, b = await _upsert(session, Shipment, d, snap, errors, seen, i)
            ins, upd = ins + a, upd + b
        except Exception as e:
            errors.append(f"строка {i}: {e}")
    return ins, upd

def _simple(model):
    async def handler(session, ws, mapping, snap, errors):
        ins = upd = 0; seen: set[str] = set()
        hdr, rows = _read(ws); idx = _mapped(hdr, mapping)
        if "order_id" not in idx:
            raise ValueError("не найдена колонка «Заказ»")
        for i, row in rows:
            try:
                d = {f: _coerce(model, f, row[j] if j < len(row) else None) for f, j in idx.items()}
                if not d.get("order_id"):
                    errors.append(f"строка {i}: пустой order_id"); continue
                if "transit_days" in model.__table__.columns and d.get("transit_days") is None \
                        and d.get("ship_date") and d.get("plan_arrival"):
                    d["transit_days"] = (d["plan_arrival"] - d["ship_date"]).days
                a, b = await _upsert(session, model, d, snap, errors, seen, i)
                ins, upd = ins + a, upd + b
            except Exception as e:
                errors.append(f"строка {i}: {e}")
        return ins, upd
    return handler

@dataclass
class Sheet:
    model: type
    handler: Callable
    default: dict

SHEETS: dict[str, Sheet] = {
    MAIN_SHEET: Sheet(Shipment, _handle_main, {}),
    "ПланНеделя": Sheet(Shipment, _handle_plan, {
        "заказ": "order_id", "адрес доставки": "address", "вес": "weight", "дата отгрузки (план)": "ship_date",
        "перевозчик": "carrier", "машина": "vehicle_raw", "водитель": "driver_raw", "таможня": "notes",
        "машина (тип)": "vehicle_type", "ставка": "rate", "прогноз": "comment"}),
    "Входящие": Sheet(InboundShipment, _simple(InboundShipment), {
        "номер недели": "week_number", "заказ": "order_id", "груз": "cargo", "город загрузки": "loading_city",
        "вес": "weight", "отправитель": "sender", "дата отгрузки (факт)": "ship_date", "дата прибытия": "plan_arrival",
        "перевозчик": "carrier", "машина": "vehicle", "водитель": "driver", "ставка": "rate", "срок транзита": "transit_days"}),
    "Местные": Sheet(LocalShipment, _simple(LocalShipment), {
        "№": "number", "no": "number", "заказ": "order_id", "груз": "cargo", "загрузка": "loading_city",
        "отправитель": "sender", "выгрузка": "unloading_city", "получатель": "receiver", "вес": "weight",
        "дата": "ship_date", "перевозчик": "carrier", "машина": "vehicle", "водитель": "driver", "ставка": "rate"}),
    "Москва": Sheet(MoscowShipment, _simple(MoscowShipment), {
        "номер недели": "week_number", "заказ": "order_id", "груз": "cargo", "город загрузки": "loading_city",
        "город выгрузки": "unloading_city", "вес": "weight", "отправитель": "sender", "получатель": "receiver",
        "дата отгрузки (факт)": "ship_date", "дата прибытия": "plan_arrival", "перевозчик": "carrier",
        "машина": "vehicle", "водитель": "driver", "ставка": "rate", "срок транзита": "transit_days"}),
    "В пути на Тэнкард": Sheet(TransitToTencard, _simple(TransitToTencard), {
        "заказ": "order_id", "город": "city", "вес": "weight", "дата отгрузки": "ship_date",
        "предполагаемая дата прихода": "plan_arrival", "перевозчик": "carrier",
        "авто/прицеп/контейнер": "vehicle", "водитель": "driver", "тип": "vehicle_type"}),
}
IGNORED = ["Контейнера КЛд-НСк", "Лист2", "погода прогноз"]

def sheet_names(src) -> list[str]:
    return load_workbook(src, read_only=True).sheetnames

def fields_of(sheet: str) -> list[str]:
    return [c.name for c in SHEETS[sheet].model.__table__.columns]

# ---------- маппинг в settings ----------
async def get_mapping(session, sheet: str) -> dict | None:
    s = await session.get(Setting, f"excel_mapping:{sheet}")
    return s.value if s else None

async def set_mapping(session, sheet: str, mapping: dict | None) -> None:
    s = await session.get(Setting, f"excel_mapping:{sheet}")
    if mapping is None:
        if s:
            await session.delete(s)
    elif s:
        s.value = mapping
    else:
        session.add(Setting(key=f"excel_mapping:{sheet}", value=mapping))
    await session.commit()

def preview(src, sheet_name: str, saved: dict | None = None) -> dict:
    ws = load_workbook(src, read_only=True)[sheet_name]
    it = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else "" for h in next(it, [])]
    rows = []
    for r in it:
        if any(c not in (None, "") for c in r):
            rows.append(["" if c is None else (_ser(c) if isinstance(c, (date, datetime)) else str(c)) for c in r])
        if len(rows) == 5:
            break
    default = SHEETS[sheet_name].default or {h.lower(): COLS[i] for i, h in enumerate(headers) if h and i < len(COLS)}
    return {"headers": headers, "rows": rows, "mapping_default": default, "mapping_saved": saved,
            "fixed_mapping": sheet_name == MAIN_SHEET, "fields": fields_of(sheet_name)}

# ---------- импорт ----------
async def import_xlsx(session: AsyncSession, src, sheet_name: str | None = None, mapping: dict | None = None, *,
                      sheets: list[str] | None = None, filename: str = "", uploaded_by: str = "",
                      dry_run: bool = False) -> dict:
    wb = load_workbook(src, read_only=True, data_only=False)
    wanted = sheets if sheets is not None else ([sheet_name] if sheet_name not in (None, "all") else list(SHEETS))
    batches = []
    for name in wanted:
        if name not in SHEETS or name not in wb.sheetnames:
            batches.append({"sheet_name": name, "inserted": 0, "updated": 0, "status": "failed",
                            "errors": [f"лист «{name}» не найден или не поддерживается"]})
            continue
        sh = SHEETS[name]
        eff = (mapping or {}).get(name) or await get_mapping(session, name) or sh.default
        snap = {"created": [], "updated": []}; errors: list[str] = []; ins = upd = 0
        try:
            ins, upd = await sh.handler(session, wb[name], eff, snap, errors)
            status = "success" if not errors else ("partial" if ins or upd else "failed")
        except Exception as e:
            await session.rollback()
            errors.append(f"лист: {e}"); status = "failed"; snap = {"created": [], "updated": []}; ins = upd = 0
        res = {"sheet_name": name, "inserted": ins, "updated": upd, "errors": errors, "status": status}
        if dry_run:
            await session.rollback()
        else:
            b = ImportBatch(filename=filename, uploaded_by=uploaded_by, sheet_name=name, inserted=ins, updated=upd,
                            errors=errors, status=status, snapshot=snap, uploaded_at=datetime.now(timezone.utc))
            session.add(b); await session.commit(); res["batch_id"] = b.id
        batches.append(res)
    return {"batches": batches, "total": {"inserted": sum(b["inserted"] for b in batches),
            "updated": sum(b["updated"] for b in batches), "errors_count": sum(len(b["errors"]) for b in batches)}}

async def rollback(session: AsyncSession, batch_id: int) -> dict:
    """Откат батча. Поля, изменённые вручную ПОСЛЕ импорта, не затираются (восстанавливаем поле только если
    оно всё ещё равно значению, записанному импортом). Созданные записи, изменённые позже, не удаляются (kept)."""
    b = await session.get(ImportBatch, batch_id)
    if b is None:
        raise LookupError("батч не найден")
    if b.status == "rolled_back":
        raise ValueError("батч уже откатен")
    model = SHEETS[b.sheet_name].model
    snap = b.snapshot or {}
    restored = removed = kept = 0
    for e in snap.get("updated", []):
        obj = await session.get(model, e["order_id"])
        if obj is None:
            continue
        changed = False
        for k, pv in e["prev"].items():
            if _ser(getattr(obj, k)) == e["new"][k]:
                setattr(obj, k, _deser(model, k, pv)); changed = True
        restored += changed
    for oid in snap.get("created", []):
        obj = await session.get(model, oid)
        if obj is None:
            continue
        ua = _aware(getattr(obj, "updated_at", None))
        if ua and ua > _aware(b.uploaded_at):
            kept += 1; continue
        if model is Shipment:
            await session.execute(delete(PositionRow).where(PositionRow.order_id == oid))
            await session.execute(delete(GpsLink).where(GpsLink.order_id == oid))
        await session.delete(obj); removed += 1
    b.status = "rolled_back"
    await session.commit()
    return {"ok": True, "restored": restored, "removed": removed, "kept": kept}
