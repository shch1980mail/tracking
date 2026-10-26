import io
from datetime import date
import httpx, pytest_asyncio
from openpyxl import Workbook
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from sqlalchemy.pool import StaticPool
import importers.excel  # noqa: F401  (регистрирует модели)
from db import get_session
from main import app
from models.shipment import Base

MAIN_H = ["№", "Неделя", "Заказ", "Направление", "Адрес", "вес", "Комм", "Дата отгрузки", "Дата прибытия", "Перевозчик",
          "Машина", "Водитель", "Ставка", "Ставка план", "Транзит", "Тип", "Термо", "Прим", "ссылка", "статус", "место"]

def make_xlsx(carrier="Лого", inbound_head="Заказ") -> bytes:
    wb = Workbook(); ws = wb.active; ws.title = "Исх.логистика"; ws.append(MAIN_H)
    ws.append([1, 2, "ТПК1", "мск", "Подольск", 20000, None, date(2026, 10, 1), date(2026, 10, 10), carrier,
               "Даф О1/39 / АЕ1", "Иванов Иван Иванович +79110000000", 1, 2, None, "авто", None, None,
               "https://maps.baltgps.ru/locator/index.html?t=abc", "в пути", None])
    ws.append([2, 2, None, "мск"])
    p = wb.create_sheet("ПланНеделя")
    p.append([None, "Заказ", "Адрес доставки", "вес", "Дата отгрузки (ПЛАН)", "Перевозчик", "Машина", "Водитель", "Таможня", "машина (тип)", "Ставка", "Прогноз"])
    p.append([None, "ТПК1", "Москва", 1, date(2026, 10, 1), carrier, "Даф О1/39 / АЕ1", "Иванов Иван Иванович", "Инмар", "термос", 5, "Прогноз 08:30 24.10"])
    p.append([None, "ТПК2", "Тверь", 1, date(2026, 10, 20), "Х", "Т1", "Петров Пётр Петрович", None, "авто", 5, None])
    i = wb.create_sheet("Входящие")
    i.append([None, "Номер недели", inbound_head, "Груз", None, "Город загрузки", "вес", "Отправитель", "Дата отгрузки (ФАКТ)", "Дата прибытия", "Перевозчик", "Машина", "Водитель", "Ставка", "Срок транзита"])
    i.append([None, 3, "ТПК18", "Какао", None, "Подольск", 25, "ТПК", date(2026, 1, 15), date(2026, 1, 29), "Конвой", "сборка", None, 6868, None])
    i.append([None, 3, "ТПК18", "Какао2", None, "Подольск", 25, "ТПК", date(2026, 1, 15), date(2026, 1, 29), "Конвой", "сборка", None, 6868, None])
    i.append([None, 3, "ТПК19", "Масло", None, "Ступино", 25, "ТПК", "не дата", None, "Конвой", "сборка", None, 1, None])
    l = wb.create_sheet("Местные")
    l.append([None, "No", "Заказ", "Груз", "Загрузка", "Отправитель", "Выгрузка", "Получатель", "вес", "Дата", "Перевозчик", "Машина", "Водитель", "Ставка"])
    l.append([None, 1, "ТПК358", "глазурь", "Светлогорск", "ТПК", "Калининград", "РХ", 2500, date(2026, 1, 20), "ИП", "Р247", "Магер", 5000])
    m = wb.create_sheet("Москва")
    m.append([None, "Номер недели", "Заказ", "Груз", None, "Город загрузки", "Город выгрузки", "вес", "Отправитель", "Получатель", "Дата отгрузки (ФАКТ)", "Дата прибытия", "Перевозчик", "Машина", "Водитель", "Ставка", "Срок транзита"])
    m.append([None, 41, "ТПК8312", "Какао", None, "Подольск", "Киров", 50, "ТПК", "ВМЗ", date(2026, 10, 8), None, "ДЛ", "сборка", None, 0, None])
    t = wb.create_sheet("В пути на Тэнкард")
    t.append(["Заказ", "Город", "вес", "дата отгрузки", "Предполагаемая дата прихода", "Перевозчик", "авто/прицеп/контейнер", "водитель", "тип"])
    t.append(["ТПК55", "Подольск", 20250, date(2026, 2, 4), date(2026, 2, 12), "Лого", "МАН", "Акатов", "ПАРОМ прицеп"])
    wb.create_sheet("Лист2")
    b = io.BytesIO(); wb.save(b); return b.getvalue()

@pytest_asyncio.fixture
async def maker():
    e = create_async_engine("sqlite+aiosqlite://", poolclass=StaticPool)
    async with e.begin() as c:
        await c.run_sync(Base.metadata.create_all)
    yield async_sessionmaker(e, expire_on_commit=False)

@pytest_asyncio.fixture
async def client(maker):
    async def ov():
        async with maker() as s:
            yield s
    app.dependency_overrides[get_session] = ov
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t", auth=("admin", "change-me")) as c:
        yield c
    app.dependency_overrides.clear()

def files(data: bytes):
    return {"file": ("a.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}

import pytest

@pytest.fixture(autouse=True)
def _no_rate_limit(monkeypatch):
    import ratelimit
    from config import settings
    monkeypatch.setattr(settings, "rate_limit_enabled", False)
    ratelimit.reset()
